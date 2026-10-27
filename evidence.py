"""채증 자료 생성: Word 보고서 / Excel"""
import io
from datetime import datetime, timezone, timedelta

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

KST = timezone(timedelta(hours=9))
FONT = "맑은 고딕"


def to_kst(iso: str) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(KST).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return iso


# ---------------------------------------------------------------- Word ------
def _font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _para(doc, text="", size=10, bold=False, align=None, space_after=4, color=None):
    # 제목도 '일반 문단 + 굵게'로만 작성 (헤딩 스타일·다음 단락과 함께 설정 없음)
    p = doc.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold, color)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def _shade(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _kv_table(doc, pairs, key_width=Cm(3.6)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in pairs:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = key_width, Cm(13)
        _shade(cells[0], "F2F2F2")
        _font(cells[0].paragraphs[0].add_run(k), 9.5, True)
        _font(cells[1].paragraphs[0].add_run(str(v)), 9.5)
    t.autofit = False
    t.columns[0].width, t.columns[1].width = key_width, Cm(13)
    return t


def build_evidence_docx(chosen: list[dict], meta: dict) -> bytes:
    info = meta["info"]
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)

    _para(doc, "온라인 악성 댓글 채증 자료", 18, True, WD_ALIGN_PARAGRAPH.CENTER, 2)
    _para(doc, "YouTube 댓글 수집 및 분류 결과", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 14, "595959")

    # 1. 수집 개요
    _para(doc, "1. 수집 개요", 12, True, space_after=6)
    _kv_table(doc, [
        ("대상 영상", info["title"]),
        ("영상 주소", info["url"]),
        ("채널", info["channel"]),
        ("영상 게시일", to_kst(info["published"])),
        ("피해자(고소인)", meta.get("victim") or "—"),
        ("수집자", meta.get("collector") or "—"),
        ("수집 일시", meta["collected_at"].strftime("%Y-%m-%d %H:%M:%S") + " (KST)"),
        ("수집 방법", "YouTube Data API v3 (commentThreads·comments) 공식 API로 자동 수집"),
        ("분류 방법", meta["method"] + " 후 검토자 확인"),
        ("수집 댓글 수", f'{meta["total"]}개 (영상 표시 댓글 수 {info["comment_count"]}개)'),
        ("증거 선별 수", f"{len(chosen)}개"),
    ])
    _para(doc, space_after=10)

    # 2. 유형별 요약
    _para(doc, "2. 유형별 요약", 12, True, space_after=6)
    cats = {}
    for c in chosen:
        for k in (c.get("category") or "기타").split(","):
            cats[k.strip()] = cats.get(k.strip(), 0) + 1
    authors = {}
    for c in chosen:
        authors[c["author"]] = authors.get(c["author"], 0) + 1
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    for i, h in enumerate(["유형", "건수"]):
        _shade(t.rows[0].cells[i], "D9D9D9")
        _font(t.rows[0].cells[i].paragraphs[0].add_run(h), 9.5, True)
    for k, v in sorted(cats.items(), key=lambda x: -x[1]):
        cells = t.add_row().cells
        _font(cells[0].paragraphs[0].add_run(k), 9.5)
        _font(cells[1].paragraphs[0].add_run(f"{v}건"), 9.5)
    _para(doc, space_after=4)
    repeat = [f"{a} ({n}건)" for a, n in sorted(authors.items(), key=lambda x: -x[1]) if n >= 2]
    _para(doc, f"작성자 수: {len(authors)}명" + (f" · 2건 이상 작성: {', '.join(repeat)}" if repeat else ""), 9.5,
          space_after=10)

    # 3. 댓글 상세
    _para(doc, "3. 악성 댓글 상세", 12, True, space_after=6)
    chosen = sorted(chosen, key=lambda c: c.get("published", ""))
    for n, c in enumerate(chosen, 1):
        _para(doc, f"증거 {n}", 10.5, True, space_after=3)
        _kv_table(doc, [
            ("작성자", c["author"]),
            ("작성자 채널", c.get("author_channel") or "—"),
            ("작성 일시", to_kst(c["published"]) + (" (수정됨: " + to_kst(c["updated"]) + ")"
                                                     if c.get("updated") and c["updated"] != c["published"] else "")),
            ("구분", "답글" if c.get("is_reply") else "댓글"),
            ("댓글 원문", c["text"]),
            ("유형 / 심각도", f'{c.get("category", "")} / {c.get("severity", "") or "—"}'),
            ("판단 근거", c.get("reason", "")),
            ("원문 링크", c["link"]),
            ("댓글 ID", c["comment_id"]),
        ])
        _para(doc, space_after=8)

    # 4. 무결성
    _para(doc, "4. 원본 데이터 무결성", 12, True, space_after=6)
    _para(doc, "수집 시점의 YouTube API 응답 원본을 JSON 파일로 별도 보관하였으며, 그 SHA-256 해시값은 아래와 같다. "
               "제출된 원본 파일의 해시값이 이와 일치하면 수집 이후 내용이 변경되지 않았음을 의미한다.", 9.5, space_after=4)
    _para(doc, meta["raw_hash"], 9, True, space_after=10)

    _para(doc, "※ 본 자료는 자동 수집·분류 결과를 검토자가 확인한 것으로, 모욕·명예훼손 해당 여부의 최종 판단은 수사기관 및 법원에 있음. "
               "댓글이 삭제될 수 있으므로 원문 링크 접속 화면 캡처를 함께 첨부할 것을 권장함.", 8.5, color="7F7F7F")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel -----
def build_excel(df: pd.DataFrame, chosen: list[dict], meta: dict) -> bytes:
    cols = {"label": "분류", "severity": "심각도", "category": "유형", "author": "작성자", "text": "댓글",
            "reason": "판단 근거", "likes": "좋아요", "published": "작성 일시(KST)", "is_reply": "답글",
            "link": "원문 링크", "author_channel": "작성자 채널", "comment_id": "댓글 ID"}
    out = df.copy()
    out["published"] = out["published"].map(to_kst)
    out = out[list(cols)].rename(columns=cols)
    sel = pd.DataFrame(chosen)
    if not sel.empty:
        sel["published"] = sel["published"].map(to_kst)
        sel = sel[list(cols)].rename(columns=cols)

    info = meta["info"]
    summary = pd.DataFrame([
        ("영상", info["title"]), ("주소", info["url"]), ("채널", info["channel"]),
        ("수집 일시(KST)", meta["collected_at"].strftime("%Y-%m-%d %H:%M:%S")),
        ("분류 방법", meta["method"]), ("수집 댓글", meta["total"]),
        ("악플 분류", int((df.label == "악플").sum())), ("증거 선별", len(chosen)),
        ("원본 SHA-256", meta["raw_hash"]),
    ], columns=["항목", "값"])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        summary.to_excel(w, sheet_name="요약", index=False)
        sel.to_excel(w, sheet_name="증거 선별", index=False)
        out.to_excel(w, sheet_name="전체 댓글", index=False)
        for ws in w.book.worksheets:
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width * 1.3), 60)
            ws.freeze_panes = "A2"
    return buf.getvalue()
