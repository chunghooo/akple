import re, io, json, hashlib, requests
from datetime import datetime, timezone, timedelta
import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Pt

KST = timezone(timedelta(hours=9))
API = "https://www.googleapis.com/youtube/v3/commentThreads"

# ---------- 키워드 (유형: (심각도, [단어])) ----------
RULES = {
    "욕설": (3, ["시발", "씨발", "ㅅㅂ", "ㅆㅂ", "좆", "ㅈ같", "개새끼", "개색기", "새끼", "병신", "ㅂㅅ", "븅신",
                 "지랄", "ㅈㄹ", "염병", "썅", "엠창", "느금", "니애미", "애미", "애비"]),
    "비속어": (1, ["존나", "졸라", "ㅈㄴ"]),
    "위협": (3, ["죽어라", "뒤져", "뒤지", "자살해", "죽여버", "찾아간다", "패버"]),
    "인신공격": (2, ["미친놈", "미친년", "또라이", "돌아이", "멍청", "찐따", "한심", "쓰레기", "못생", "정신병",
                     "정병", "관종", "역겹", "꺼져", "닥쳐", "ㄲㅈ", "ㄷㅊ"]),
    "혐오": (3, ["한남충", "한녀", "김치녀", "틀딱", "급식충", "맘충", "짱깨", "쪽바리"]),
    "성적": (3, ["따먹", "몸매보"]),
}
SEP = r"[\s\d._*~!^♡♥,·ㅇㅋㅎㅡ-]*"  # 끼워넣은 숫자·기호·ㅇㅋㅎ 건너뜀


def pat(w):
    return re.compile(SEP.join(map(re.escape, w)))


PATS = [(c, s, w, pat(w)) for c, (s, ws) in RULES.items() for w in ws]

# ---------- 우회 표기 복원 ----------
CHO, JUNG = "ㄱㄲㄴㄷㄸㄹㅁㅂㅃㅅㅆㅇㅈㅉㅊㅋㅌㅍㅎ", "ㅏㅐㅑㅒㅓㅔㅕㅖㅗㅘㅙㅚㅛㅜㅝㅞㅟㅠㅡㅢㅣ"
JONG = " ㄱㄲㄳㄴㄵㄶㄷㄹㄺㄻㄼㄽㄾㄿㅀㅁㅂㅄㅅㅆㅇㅈㅊㅋㅌㅍㅎ"
KEY = dict(zip("qwertasdfgzxcvyuiophjklbnmQWERTOP", "ㅂㅈㄷㄱㅅㅁㄴㅇㄹㅎㅋㅌㅊㅍㅛㅕㅑㅐㅔㅗㅓㅏㅣㅠㅜㅡㅃㅉㄸㄲㅆㅒㅖ"))
VC = {"ㅗㅏ": "ㅘ", "ㅗㅐ": "ㅙ", "ㅗㅣ": "ㅚ", "ㅜㅓ": "ㅝ", "ㅜㅔ": "ㅞ", "ㅜㅣ": "ㅟ", "ㅡㅣ": "ㅢ"}
ROMAN = {"ssi": "씨", "shi": "시", "si": "시", "c": "시", "bal": "발", "ppal": "빨", "byung": "병", "byeong": "병",
         "shin": "신", "sin": "신", "gae": "개", "sae": "새", "kki": "끼", "ji": "지", "ral": "랄", "jot": "좆"}
TENSE = dict(zip("ㄲㄸㅃㅆㅉ", "ㄱㄷㅂㅅㅈ"))
VOW = dict(zip("ㅢㅟㅐㅙㅞㅚㅑㅘㅕㅝㅛㅠ", "ㅣㅣㅔㅔㅔㅔㅏㅏㅓㅓㅗㅜ"))


def syl(c, v, j=" "):
    return chr(0xAC00 + (CHO.index(c) * 21 + JUNG.index(v)) * 28 + JONG.index(j))


def eng2kor(w):
    """영타 → 한글. 받침 하나만 처리(욕설 복원엔 충분). 낱자가 남으면 실패(None)"""
    j = [KEY.get(ch) or KEY.get(ch.lower()) for ch in w]
    if None in j:
        return None
    out, i = "", 0
    while i < len(j):
        if i + 1 < len(j) and j[i] in CHO and j[i + 1] in JUNG:
            c, v, i = j[i], j[i + 1], i + 2
            if i < len(j) and j[i] in JUNG and v + j[i] in VC:
                v, i = VC[v + j[i]], i + 1
            t = " "
            if i < len(j) and j[i] in JONG and not (i + 1 < len(j) and j[i + 1] in JUNG):
                t, i = j[i], i + 1
            out += syl(c, v, t)
        else:
            return None
    return out


def roman2kor(w):
    w, out = w.lower(), ""
    while w:
        k = next((k for k in sorted(ROMAN, key=len, reverse=True) if w.startswith(k)), None)
        if not k:
            return None
        out, w = out + ROMAN[k], w[len(k):]
    return out


def deob(t):
    """영어 문장 댓글은 건드리지 않음(영어 단어가 엉뚱한 욕으로 바뀌는 것 방지)"""
    words = re.findall(r"[A-Za-z]+", t)
    if len(words) >= 4 and len(re.findall(r"[가-힣ㄱ-ㅣ]", t)) <= len(t) * 0.1:
        return t
    return re.sub(r"[A-Za-z]+", lambda m: roman2kor(m[0]) or eng2kor(m[0]) or m[0], t)


def phon(t):
    """된소리→예사소리, 비슷한 모음 통일 (씌발·쉬발·씨빨 → 시발)"""
    out = ""
    for ch in t:
        n = ord(ch) - 0xAC00
        if 0 <= n < 11172:
            c, v, j = CHO[n // 588], JUNG[n % 588 // 28], JONG[n % 28]
            ch = syl(TENSE.get(c, c), VOW.get(v, v), j)
        out += ch
    return out


CORE = [(w, pat(phon(w))) for w in ["시발", "병신", "븅신", "지랄", "염병", "개새끼", "니애미"]]


def classify(text):
    d = deob(text)
    hits = []
    for c, s, w, p in PATS:
        if p.search(text):
            hits.append((s, c, f"키워드 '{w}'"))
        elif d != text and re.search("[가-힣]", w) and p.search(d):
            hits.append((s, c, f"우회 표기 '{w}'"))
    if not hits:
        ph = phon(d)
        hits = [(3, "욕설", f"발음 변형 '{w}'") for w, p in CORE if p.search(ph)][:1]
    if not hits:
        return dict(label="정상", severity=0, category="", reason="")
    hits.sort(reverse=True)
    return dict(label="악플", severity=hits[0][0], category=", ".join(dict.fromkeys(h[1] for h in hits)),
                reason=", ".join(dict.fromkeys(h[2] for h in hits)))


# ---------- 유튜브 수집 ----------
def video_id(url):
    m = re.search(r"(?:v=|youtu\.be/|shorts/|embed/|live/)([\w-]{11})", url)
    return m[1] if m else None


def row(c, vid):
    s = c["snippet"]
    return dict(id=c["id"], author=s["authorDisplayName"], channel=s.get("authorChannelUrl", ""),
                text=s["textOriginal"], time=s["publishedAt"], link=f"https://youtu.be/{vid}?lc={c['id']}")


def fetch(vid, key, limit):
    rows, raw, token = [], [], None
    while len(rows) < limit:
        r = requests.get(API, params=dict(part="snippet,replies", videoId=vid, maxResults=100,
                                          textFormat="plainText", key=key, pageToken=token or ""), timeout=20)
        if r.status_code != 200:
            raise RuntimeError(r.json().get("error", {}).get("message", r.text[:200]))
        d = r.json()
        raw.append(d)
        for th in d["items"]:
            rows.append(row(th["snippet"]["topLevelComment"], vid))
            rows += [row(c, vid) for c in th.get("replies", {}).get("comments", [])]
        token = d.get("nextPageToken")
        if not token:
            break
    return rows[:limit], raw


# ---------- 채증 보고서 ----------
def kst(t):
    return datetime.fromisoformat(t.replace("Z", "+00:00")).astimezone(KST).strftime("%Y-%m-%d %H:%M")


def report(sel, url, at, total, sha):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)

    def p(t, size=10, bold=False):
        r = doc.add_paragraph().add_run(t)
        r.bold, r.font.size = bold, Pt(size)

    def table(pairs):
        tb = doc.add_table(0, 2, style="Table Grid")
        for k, v in pairs:
            a, b = tb.add_row().cells
            a.text, b.text = k, str(v)
        doc.add_paragraph()

    p("온라인 악성 댓글 채증 자료", 16, True)
    p("1. 수집 개요", 12, True)
    table([("영상 주소", url), ("수집 일시", at + " (KST)"), ("수집 방법", "YouTube Data API v3 공식 API"),
           ("수집 댓글 수", total), ("증거 선별 수", len(sel))])
    p("2. 악성 댓글 상세", 12, True)
    for i, c in enumerate(sel, 1):
        p(f"증거 {i}", 10, True)
        table([("작성자", c["author"]), ("작성자 채널", c["channel"]), ("작성 일시", kst(c["time"])),
               ("댓글 원문", c["text"]), ("유형/심각도", f'{c["category"]} / {c["severity"]}'),
               ("판단 근거", c["reason"]), ("원문 링크", c["link"])])
    p("3. 원본 데이터 무결성", 12, True)
    p(f"원본 JSON SHA-256: {sha}")
    p("※ 최종 판단은 수사기관·법원에 있으며, 원문 링크 화면 캡처를 함께 첨부할 것을 권장함.", 9)
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


# ---------- 화면 ----------
st.set_page_config("악플 수집기", "🛡️")
st.title("🛡️ 악플 수집기")
try:
    key = st.secrets["YOUTUBE_API_KEY"]
except Exception:
    key = st.text_input("YouTube API 키", type="password")
url = st.text_input("유튜브 영상 주소")
limit = st.slider("최대 수집 댓글 수", 100, 3000, 1000, 100)

if st.button("수집 시작", type="primary"):
    vid = video_id(url)
    if not (vid and key):
        st.error("주소와 API 키를 확인하세요.")
        st.stop()
    with st.spinner("수집·분류 중…"):
        try:
            rows, raw = fetch(vid, key, limit)
        except Exception as e:
            st.error(f"수집 실패: {e}")
            st.stop()
        for r in rows:
            r.update(classify(r["text"]))
        rawb = json.dumps(raw, ensure_ascii=False).encode()
        st.session_state.d = dict(rows=rows, raw=rawb, sha=hashlib.sha256(rawb).hexdigest(),
                                  at=datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S"), url=url, vid=vid)

if "d" in st.session_state:
    d = st.session_state.d
    df = pd.DataFrame(d["rows"])
    st.metric("악플 / 전체", f"{(df.label == '악플').sum()} / {len(df)}")

    # 필터: 보고 싶은 단어·유형만 골라보기 (악플이 아닌 댓글도 검색 가능)
    f1, f2 = st.columns(2)
    scope = f1.radio("범위", ["악플만", "전체 댓글"], horizontal=True)
    cats = f2.multiselect("유형", list(RULES))
    q = st.text_input("검색어 (쉼표로 여러 개, 띄어쓰기·숫자 끼운 변형도 찾음)")
    v = df if scope == "전체 댓글" else df[df.label == "악플"]
    if cats:
        v = v[v.category.apply(lambda c: any(x in c for x in cats))]
    words = [w.strip() for w in q.split(",") if w.strip()]
    if words:
        v = v[v.text.apply(lambda t: any(pat(w).search(t) for w in words))]
    v = v.sort_values("severity", ascending=False).copy()
    v.insert(0, "포함", v.label == "악플")
    st.caption(f"{len(v)}개 표시 · 증거에 넣을 댓글을 체크하세요")
    ed = st.data_editor(v[["포함", "severity", "category", "author", "text", "reason", "link"]], hide_index=True,
                        column_config={"link": st.column_config.LinkColumn(display_text="열기")})
    sel = df.loc[ed[ed["포함"]].index].to_dict("records")
    for c in sel:
        if c["label"] == "정상":
            c.update(category="사용자 선택", reason="검토자가 직접 선택" + (f" (검색어: {q})" if q else ""))
    c1, c2, c3 = st.columns(3)
    c1.download_button("채증 보고서(Word)", report(sel, d["url"], d["at"], len(df), d["sha"]),
                       f"채증_{d['vid']}.docx", disabled=not sel)
    c2.download_button("전체 댓글(CSV)", df.to_csv(index=False).encode("utf-8-sig"), f"댓글_{d['vid']}.csv")
    c3.download_button("원본(JSON)", d["raw"], f"원본_{d['vid']}.json")
    st.caption(f"원본 SHA-256: {d['sha']}")
